"""Genera il prodotto "Gestionale Forfettario" (xlsx) da caricare su Lemon Squeezy.

Uso: python3 prodotti/genera_gestionale.py [percorso_output.xlsx]
Il file generato NON va messo nel sito né nel repository (è il prodotto a pagamento).
"""
import sys

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, DataBarRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

VERDE = "1F6F5C"
CHIARO = "E8F3EF"
EURO = '#,##0.00 "€"'
PERC = "0.00%"
H = Font(bold=True, color="FFFFFF")
FILL_H = PatternFill("solid", fgColor=VERDE)
FILL_IN = PatternFill("solid", fgColor="FFF6E0")
FILL_OUT = PatternFill("solid", fgColor=CHIARO)
MESI = ["Gennaio", "Febbraio", "Marzo", "Aprile", "Maggio", "Giugno", "Luglio",
        "Agosto", "Settembre", "Ottobre", "Novembre", "Dicembre"]
RIGHE = 500


def header(ws, row, values):
    for col, v in enumerate(values, 1):
        c = ws.cell(row=row, column=col, value=v)
        c.font, c.fill = H, FILL_H
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def istruzioni(wb):
    ws = wb.active
    ws.title = "Istruzioni"
    ws.column_dimensions["A"].width = 100
    righe = [
        ("GESTIONALE FORFETTARIO – ContiChiari", Font(bold=True, size=16, color=VERDE)),
        ("", None),
        ("Come si usa, in 3 passi:", Font(bold=True, size=12)),
        ("1. Vai nel foglio «Impostazioni» e compila le celle gialle: anno, coefficiente, imposta (5% o 15%) e gestione INPS.", None),
        ("2. Nel foglio «Incassi» registra ogni fattura quando viene PAGATA (nel forfettario conta la data di incasso).", None),
        ("3. Apri il foglio «Riepilogo» per vedere tasse, contributi, netto, accantonamento mensile e distanza dal limite di 85.000 €.", None),
        ("", None),
        ("Legenda colori:", Font(bold=True, size=12)),
        ("Celle gialle = da compilare. Celle verdi = calcolate in automatico, non modificarle.", None),
        ("", None),
        ("Note importanti:", Font(bold=True, size=12)),
        ("• Le aliquote INPS e i contributi fissi cambiano ogni anno: aggiornali nelle Impostazioni con i valori ufficiali INPS.", None),
        ("• Il calcolo è una stima semplificata (considera deducibili nello stesso anno i contributi calcolati) e non include acconti, saldi e marche da bollo.", None),
        ("• Non sostituisce il commercialista: serve a tenere sotto controllo i tuoi numeri durante l'anno.", None),
        ("• Funziona con Excel, Google Sheets (File > Importa), LibreOffice e Numbers.", None),
    ]
    for i, (t, f) in enumerate(righe, 1):
        c = ws.cell(row=i, column=1, value=t)
        c.alignment = Alignment(wrap_text=True)
        if f:
            c.font = f


def impostazioni(wb):
    ws = wb.create_sheet("Impostazioni")
    ws.column_dimensions["A"].width = 48
    ws.column_dimensions["B"].width = 22
    ws.column_dimensions["C"].width = 60
    header(ws, 1, ["Parametro", "Valore", "Note"])
    dati = [
        ("Anno", 2026, "Anno fiscale di riferimento", None),
        ("Coefficiente di redditività", 0.78, "78% professionisti, 67% altre attività, 62% intermediari, 86% costruzioni, 54% ambulanti non alimentari, 40% commercio/ristorazione", PERC),
        ("Imposta sostitutiva", 0.15, "15% ordinaria, 5% per i primi 5 anni se hai i requisiti", PERC),
        ("Gestione INPS", "Gestione Separata", "Gestione Separata / Artigiani / Commercianti / Cassa professionale", None),
        ("Aliquota Gestione Separata o Cassa", 0.2607, "Solo per Gestione Separata o Cassa: aggiorna ogni anno", PERC),
        ("Contributi fissi annui (Artigiani/Commercianti)", 4460, "Solo per Artigiani/Commercianti: valore indicativo, verifica sul sito INPS", EURO),
        ("Reddito minimale (Artigiani/Commercianti)", 18555, "Solo per Artigiani/Commercianti: valore indicativo", EURO),
        ("Aliquota oltre il minimale", 0.24, "Artigiani circa 24%, Commercianti circa 24,48%", PERC),
        ("Riduzione contributiva 35%", "No", "Sì / No – solo Artigiani/Commercianti che l'hanno richiesta", None),
        ("Limite ricavi forfettario", 85000, "Soglia per restare nel regime l'anno successivo", EURO),
    ]
    for i, (p, v, n, fmt) in enumerate(dati, 2):
        ws.cell(row=i, column=1, value=p).font = Font(bold=True)
        c = ws.cell(row=i, column=2, value=v)
        c.fill = FILL_IN
        if fmt:
            c.number_format = fmt
        ws.cell(row=i, column=3, value=n).alignment = Alignment(wrap_text=True)
    dv = DataValidation(type="list", formula1='"Gestione Separata,Artigiani,Commercianti,Cassa professionale"', allow_blank=False)
    ws.add_data_validation(dv)
    dv.add("B5")
    dv2 = DataValidation(type="list", formula1='"Sì,No"', allow_blank=False)
    ws.add_data_validation(dv2)
    dv2.add("B10")
    dv3 = DataValidation(type="list", formula1='"0.15,0.05"', allow_blank=False)
    ws.add_data_validation(dv3)
    dv3.add("B4")


def incassi(wb):
    ws = wb.create_sheet("Incassi")
    widths = [14, 12, 32, 40, 16, 8, 8]
    for i, w in enumerate(widths):
        ws.column_dimensions["ABCDEFG"[i]].width = w
    header(ws, 1, ["Data incasso", "N. fattura", "Cliente", "Descrizione", "Importo incassato", "Mese", "Anno"])
    ws.freeze_panes = "A2"
    for r in range(2, RIGHE + 2):
        ws.cell(row=r, column=1).number_format = "DD/MM/YYYY"
        ws.cell(row=r, column=5).number_format = EURO
        for col in range(1, 6):
            ws.cell(row=r, column=col).fill = FILL_IN
        m = ws.cell(row=r, column=6, value=f'=IF(A{r}="","",MONTH(A{r}))')
        m.fill = FILL_OUT
        y = ws.cell(row=r, column=7, value=f'=IF(A{r}="","",YEAR(A{r}))')
        y.fill = FILL_OUT
    dv = DataValidation(type="date", operator="greaterThan", formula1="DATE(2000,1,1)")
    dv.error, dv.errorTitle = "Inserisci una data valida (gg/mm/aaaa)", "Data non valida"
    ws.add_data_validation(dv)
    dv.add(f"A2:A{RIGHE + 1}")


def riepilogo(wb):
    ws = wb.create_sheet("Riepilogo")
    ws.column_dimensions["A"].width = 46
    ws.column_dimensions["B"].width = 20
    inc = f"Incassi!$E$2:$E${RIGHE + 1}"
    mesi = f"Incassi!$F$2:$F${RIGHE + 1}"
    anni = f"Incassi!$G$2:$G${RIGHE + 1}"
    righe = [
        ("Ricavi incassati nell'anno", f"=SUMIFS({inc},{anni},Impostazioni!$B$2)", EURO),
        ("Reddito lordo (ricavi × coefficiente)", "=B2*Impostazioni!$B$3", EURO),
        ("Contributi INPS stimati",
         '=IF(OR(Impostazioni!$B$5="Gestione Separata",Impostazioni!$B$5="Cassa professionale"),'
         "B3*Impostazioni!$B$6,"
         "(Impostazioni!$B$7+MAX(0,B3-Impostazioni!$B$8)*Impostazioni!$B$9)"
         '*IF(Impostazioni!$B$10="Sì",0.65,1))', EURO),
        ("Base imponibile (reddito − contributi)", "=MAX(0,B3-B4)", EURO),
        ("Imposta sostitutiva stimata", "=B5*Impostazioni!$B$4", EURO),
        ("Totale tasse + contributi", "=B4+B6", EURO),
        ("Netto stimato", "=B2-B7", EURO),
        ("Pressione su ricavi", "=IF(B2=0,0,B7/B2)", PERC),
        ("Percentuale da accantonare su ogni incasso", "=B9", PERC),
        ("Limite forfettario usato", "=B2/Impostazioni!$B$11", PERC),
        ("Ricavi ancora disponibili prima del limite", "=MAX(0,Impostazioni!$B$11-B2)", EURO),
    ]
    header(ws, 1, ["Riepilogo annuale", "Valore"])
    for i, (label, f, fmt) in enumerate(righe, 2):
        ws.cell(row=i, column=1, value=label).font = Font(bold=True)
        c = ws.cell(row=i, column=2, value=f)
        c.number_format, c.fill = fmt, FILL_OUT
    ws["B8"].font = Font(bold=True, size=13, color=VERDE)
    ws.conditional_formatting.add("B11", DataBarRule(start_type="num", start_value=0, end_type="num", end_value=1, color="1F6F5C"))
    ws.conditional_formatting.add("B11", CellIsRule(operator="greaterThan", formula=["0.9"], font=Font(bold=True, color="C00000")))

    start = 15
    ws.cell(row=start - 1, column=1, value="Andamento mensile").font = Font(bold=True, size=13, color=VERDE)
    header(ws, start, ["Mese", "Incassato", "Da accantonare", "Netto stimato", "Cumulato anno"])
    for col in "CDE":
        ws.column_dimensions[col].width = 18
    for i, mese in enumerate(MESI, 1):
        r = start + i
        ws.cell(row=r, column=1, value=mese)
        ws.cell(row=r, column=2, value=f"=SUMIFS({inc},{anni},Impostazioni!$B$2,{mesi},{i})")
        ws.cell(row=r, column=3, value=f"=B{r}*$B$9")
        ws.cell(row=r, column=4, value=f"=B{r}-C{r}")
        ws.cell(row=r, column=5, value=f"=SUM($B${start + 1}:B{r})")
        for col in range(2, 6):
            c = ws.cell(row=r, column=col)
            c.number_format, c.fill = EURO, FILL_OUT


def main():
    out = sys.argv[1] if len(sys.argv) > 1 else "Gestionale-Forfettario-ContiChiari.xlsx"
    wb = Workbook()
    istruzioni(wb)
    impostazioni(wb)
    incassi(wb)
    riepilogo(wb)
    wb.save(out)
    print("Creato", out)


if __name__ == "__main__":
    main()
